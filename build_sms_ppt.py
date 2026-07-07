#!/usr/bin/env python3
"""
Safety Management System (SMS) - Professional presentation builder.

Generates a polished, 16:9 PowerPoint deck with a consistent aviation theme,
icon badges and slide entrance animations (fade / fly-in) injected as raw XML.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
import copy

# ----------------------------------------------------------------------------
# Theme
# ----------------------------------------------------------------------------
NAVY    = RGBColor(0x0B, 0x25, 0x45)   # deep background
DEEP    = RGBColor(0x13, 0x35, 0x66)   # panel
SKY     = RGBColor(0x2E, 0x86, 0xDE)   # primary accent
SKY_LT  = RGBColor(0x6F, 0xB1, 0xF0)
TEAL    = RGBColor(0x1A, 0xBC, 0x9C)   # secondary accent
GOLD    = RGBColor(0xF4, 0xB7, 0x40)   # highlight
LIGHT   = RGBColor(0xF4, 0xF7, 0xFB)   # light background
CARD    = RGBColor(0xFF, 0xFF, 0xFF)
INK     = RGBColor(0x1B, 0x26, 0x38)   # dark text
GRAY    = RGBColor(0x5A, 0x64, 0x72)   # muted text
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SHADOW  = RGBColor(0xD5, 0xDE, 0xEA)

HEAD_FONT = "Calibri"
BODY_FONT = "Calibri"

EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------
# Low level helpers
# ----------------------------------------------------------------------------
def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _no_line(shape):
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color, shape_type=MSO_SHAPE.RECTANGLE, line=None, line_w=None):
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    sp.shadow.inherit = False
    _set_fill(sp, color)
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    return sp


def add_gradient_bg(slide, c1, c2, angle=90):
    """Full-slide two-stop linear gradient."""
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    sp.shadow.inherit = False
    sp.line.fill.background()
    fill = sp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = c1
    stops[1].position = 1.0
    stops[1].color.rgb = c2
    try:
        fill.gradient_angle = angle
    except Exception:
        pass
    return sp


def add_solid_bg(slide, color):
    return add_rect(slide, 0, 0, EMU_W, EMU_H, color)


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.0, space_after=Pt(6), wrap=True):
    """runs: list of paragraphs; each paragraph is a list of (text, dict) run specs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = space_after
        p.space_before = Pt(0)
        for text, spec in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = spec.get("font", BODY_FONT)
            f.size = spec.get("size", Pt(18))
            f.bold = spec.get("bold", False)
            f.italic = spec.get("italic", False)
            f.color.rgb = spec.get("color", INK)
    return tb


def add_icon_badge(slide, x, y, d, color, glyph, glyph_color=WHITE, glyph_size=None):
    """Round colored badge with a unicode/emoji glyph."""
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    circ.shadow.inherit = False
    _set_fill(circ, color)
    tf = circ.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.size = glyph_size or Pt(int(d / 22000))
    r.font.color.rgb = glyph_color
    r.font.name = "Segoe UI Emoji"
    return circ


def footer(slide, idx):
    add_rect(slide, 0, EMU_H - Inches(0.14), EMU_W, Inches(0.14), SKY)
    add_rect(slide, 0, EMU_H - Inches(0.14), Inches(3.2), Inches(0.14), TEAL)
    add_text(slide, EMU_W - Inches(4.2), EMU_H - Inches(0.52), Inches(3.6), Inches(0.3),
             [[("Safety Management System   |   %02d" % idx, {"size": Pt(9), "color": GRAY})]],
             align=PP_ALIGN.RIGHT, wrap=False)


def section_header(slide, kicker, title, num):
    """Standard light content-slide header with accent bar."""
    add_rect(slide, Inches(0.6), Inches(0.62), Inches(0.16), Inches(0.9), GOLD)
    add_text(slide, Inches(0.95), Inches(0.55), Inches(10.5), Inches(0.35),
             [[(kicker, {"size": Pt(13), "bold": True, "color": SKY, "font": HEAD_FONT})]])
    add_text(slide, Inches(0.95), Inches(0.85), Inches(11.4), Inches(0.8),
             [[(title, {"size": Pt(30), "bold": True, "color": NAVY, "font": HEAD_FONT})]])


# ----------------------------------------------------------------------------
# Animation injection (entrance effects, auto play after previous)
# ----------------------------------------------------------------------------
class AnimBuilder:
    def __init__(self):
        self._id = 3  # 1 & 2 reserved for tmRoot/mainSeq
        self.pars = []

    def _nid(self):
        v = self._id
        self._id += 1
        return v

    def add(self, spid, effect="fade", delay=0, dur=500):
        """effect: 'fade' | 'flyLeft' | 'flyBottom' | 'zoom'."""
        outer = self._nid()
        inner = self._nid()
        eff = self._nid()
        setn = self._nid()
        anim = self._nid()

        if effect == "flyLeft":
            preset_id, subtype = 2, 4
            move = ('<p:anim calcmode="lin" valueType="num">'
                    '<p:cBhvr additive="base"><p:cTn id="%d" dur="%d" fill="hold"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
                    '<p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr>'
                    '<p:tavLst><p:tav tm="0"><p:val><p:strVal val="0-#ppt_w/2"/></p:val></p:tav>'
                    '<p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim>'
                    % (anim, dur, spid))
        elif effect == "flyBottom":
            preset_id, subtype = 2, 8
            move = ('<p:anim calcmode="lin" valueType="num">'
                    '<p:cBhvr additive="base"><p:cTn id="%d" dur="%d" fill="hold"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
                    '<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                    '<p:tavLst><p:tav tm="0"><p:val><p:strVal val="1+#ppt_h/2"/></p:val></p:tav>'
                    '<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'
                    % (anim, dur, spid))
        elif effect == "zoom":
            preset_id, subtype = 23, 0
            move = ('<p:animEffect transition="in" filter="fade">'
                    '<p:cBhvr><p:cTn id="%d" dur="%d"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    '<p:anim calcmode="lin" valueType="num">'
                    '<p:cBhvr additive="base"><p:cTn id="%d" dur="%d" fill="hold"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
                    '<p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr>'
                    '<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0.4"/></p:val></p:tav>'
                    '<p:tav tm="100000"><p:val><p:fltVal val="1"/></p:val></p:tav></p:tavLst></p:anim>'
                    '<p:anim calcmode="lin" valueType="num">'
                    '<p:cBhvr additive="base"><p:cTn id="%d" dur="%d" fill="hold"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
                    '<p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr>'
                    '<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0.4"/></p:val></p:tav>'
                    '<p:tav tm="100000"><p:val><p:fltVal val="1"/></p:val></p:tav></p:tavLst></p:anim>'
                    % (anim, dur, spid, self._nid(), dur, spid, self._nid(), dur, spid))
        else:  # fade
            preset_id, subtype = 10, 0
            move = ('<p:animEffect transition="in" filter="fade">'
                    '<p:cBhvr><p:cTn id="%d" dur="%d"/>'
                    '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    % (anim, dur, spid))

        par = (
            '<p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:cTn id="%d" fill="hold">'
            '<p:stCondLst><p:cond delay="%d"/></p:stCondLst>'
            '<p:childTnLst><p:par><p:cTn id="%d" fill="hold">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst><p:par><p:cTn id="%d" presetID="%d" presetClass="entr" '
            'presetSubtype="%d" fill="hold" grpId="0" nodeType="afterEffect">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst>'
            '<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
            '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            '<p:to><p:strVal val="visible"/></p:to></p:set>'
            '%s'
            '</p:childTnLst></p:cTn></p:par></p:childTnLst>'
            '</p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
            % (outer, delay, inner, eff, preset_id, subtype, setn, spid, move)
        )
        self.pars.append(par)

    def xml(self):
        if not self.pars:
            return None
        body = "".join(self.pars)
        return (
            '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
            '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            + body +
            '</p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
        )


from lxml import etree


def apply_anim(slide, builder):
    xml = builder.xml()
    if not xml:
        return
    el = etree.fromstring(xml)
    slide.shapes._spTree.getparent().append(el)


def sid(shape):
    return shape.shape_id


# ============================================================================
# SLIDE 1 — TITLE
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s, NAVY, DEEP, angle=45)
# decorative circles
c1 = add_rect(s, Inches(9.4), Inches(-1.6), Inches(5), Inches(5), DEEP, MSO_SHAPE.OVAL)
c1.fill.fore_color.rgb = RGBColor(0x18, 0x3E, 0x74)
c2 = add_rect(s, Inches(10.8), Inches(4.4), Inches(4), Inches(4), DEEP, MSO_SHAPE.OVAL)
c2.fill.fore_color.rgb = RGBColor(0x16, 0x39, 0x6B)
# thin accent lines
add_rect(s, Inches(1.0), Inches(2.35), Inches(1.3), Inches(0.06), GOLD)
badge = add_icon_badge(s, Inches(1.0), Inches(1.05), Inches(1.05), TEAL, "\u2708",
                       glyph_size=Pt(30))
t_title = add_text(s, Inches(1.0), Inches(2.7), Inches(10.5), Inches(2.2),
                   [[("SAFETY MANAGEMENT", {"size": Pt(54), "bold": True, "color": WHITE, "font": HEAD_FONT})],
                    [("SYSTEM", {"size": Pt(54), "bold": True, "color": GOLD, "font": HEAD_FONT})]],
                   line_spacing=1.02, space_after=Pt(0))
t_sub = add_text(s, Inches(1.05), Inches(4.85), Inches(10), Inches(0.6),
                 [[("A systematic, proactive approach to managing safety in aviation",
                    {"size": Pt(18), "color": SKY_LT, "italic": True})]])
t_meta = add_text(s, Inches(1.05), Inches(6.15), Inches(10), Inches(0.6),
                  [[("Aligned with  ", {"size": Pt(13), "color": RGBColor(0xB8,0xC7,0xDB)}),
                    ("ICAO Annex 19", {"size": Pt(13), "bold": True, "color": WHITE}),
                    ("  |  ", {"size": Pt(13), "color": RGBColor(0xB8,0xC7,0xDB)}),
                    ("DGCA CAR", {"size": Pt(13), "bold": True, "color": WHITE})]])
add_rect(s, 0, EMU_H - Inches(0.14), EMU_W, Inches(0.14), GOLD)

ab = AnimBuilder()
ab.add(sid(badge), "zoom", delay=0)
ab.add(sid(t_title), "flyLeft", delay=200)
ab.add(sid(t_sub), "fade", delay=300)
ab.add(sid(t_meta), "fade", delay=300)
apply_anim(s, ab)


# ============================================================================
# SLIDE 2 — WHAT IS SMS
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_solid_bg(s, LIGHT)
add_rect(s, 0, 0, Inches(4.6), EMU_H, NAVY)
add_rect(s, Inches(4.6), 0, Inches(0.08), EMU_H, GOLD)
ic = add_icon_badge(s, Inches(0.7), Inches(1.0), Inches(1.15), TEAL, "\U0001F6E1", glyph_size=Pt(34))
lt = add_text(s, Inches(0.7), Inches(2.6), Inches(3.5), Inches(2.2),
              [[("WHAT IS", {"size": Pt(24), "bold": True, "color": SKY_LT})],
               [("SAFETY", {"size": Pt(30), "bold": True, "color": WHITE})],
               [("MANAGEMENT", {"size": Pt(30), "bold": True, "color": WHITE})],
               [("SYSTEM?", {"size": Pt(30), "bold": True, "color": GOLD})]],
              line_spacing=1.05, space_after=Pt(0))
card = add_rect(s, Inches(5.4), Inches(2.15), Inches(7.1), Inches(3.2), CARD,
                MSO_SHAPE.ROUNDED_RECTANGLE)
add_rect(s, Inches(5.4), Inches(2.15), Inches(0.14), Inches(3.2), SKY)
dq = add_text(s, Inches(5.75), Inches(2.35), Inches(1), Inches(0.7),
              [[("\u201C", {"size": Pt(60), "bold": True, "color": SHADOW})]])
body = add_text(s, Inches(5.9), Inches(2.85), Inches(6.2), Inches(2.3),
                [[("A systematic approach to managing safety", {"size": Pt(22), "bold": True, "color": NAVY})],
                 [("", {"size": Pt(6)})],
                 [("including the necessary organizational structures, "
                   "accountabilities, responsibilities, policies and procedures.",
                   {"size": Pt(19), "color": GRAY})]],
                line_spacing=1.15)
footer(s, 2)

ab = AnimBuilder()
ab.add(sid(ic), "zoom")
ab.add(sid(lt), "flyLeft", delay=200)
ab.add(sid(card), "fade", delay=200)
ab.add(sid(dq), "fade", delay=100)
ab.add(sid(body), "flyBottom", delay=200)
apply_anim(s, ab)


# ============================================================================
# SLIDE 3 — WHY SMS REQUIRED
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_solid_bg(s, LIGHT)
section_header(s, "THE NEED FOR SMS", "Why is SMS Required in Aviation?", 3)

points = [
    ("\u2708", "Aviation is a High-Risk + High-Consequence industry."),
    ("\u26A0", "A single mistake can lead to a huge number of fatalities."),
    ("\U0001F501", "The old \u201CFind & Fix after accident\u201D model is no longer enough for such a high-risk industry."),
    ("\U0001F4CA", "SMS works by: Collect data \u2192 Analyze risk \u2192 Fix the system before fatal accidents."),
    ("\U0001F50D", "SMS forces organizations to find hazards before they become accidents."),
    ("\u2699", "In modern aviation safety can\u2019t be left to chance \u2014 it must be managed systematically, just like finance or quality."),
]
ab = AnimBuilder()
top = Inches(1.85)
row_h = Inches(0.82)
for i, (g, txt) in enumerate(points):
    y = top + i * row_h
    card = add_rect(s, Inches(0.95), y, Inches(11.45), Inches(0.72), CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
    bdg = add_icon_badge(s, Inches(1.12), y + Inches(0.11), Inches(0.5), SKY, g, glyph_size=Pt(15))
    tx = add_text(s, Inches(1.85), y, Inches(10.4), Inches(0.72),
                  [[(txt, {"size": Pt(16.5), "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    ab.add(sid(card), "flyLeft", delay=(120 if i else 0))
    ab.add(sid(bdg), "zoom", delay=0)
    ab.add(sid(tx), "fade", delay=0)
footer(s, 3)
apply_anim(s, ab)


# ============================================================================
# SLIDE 4 — HOW TO MANAGE SYSTEMATICALLY + 4 COMPONENTS
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_solid_bg(s, LIGHT)
section_header(s, "THE APPROACH", "How Will You Manage SMS Systematically?", 4)

intro = add_text(s, Inches(0.95), Inches(1.75), Inches(11.5), Inches(1.5),
                 [[("A structured, repeatable process to identify risks, manage them and keep "
                    "improving \u2014 instead of just reacting after something goes wrong.",
                    {"size": Pt(17), "color": INK})],
                  [("The SMS framework incorporates ", {"size": Pt(17), "color": INK}),
                   ("4 Components", {"size": Pt(17), "bold": True, "color": SKY}),
                   (" and ", {"size": Pt(17), "color": INK}),
                   ("12 Key Elements.", {"size": Pt(17), "bold": True, "color": TEAL})]],
                 line_spacing=1.15, space_after=Pt(8))

comps = [
    ("1", "\U0001F4CB", "Safety Policy\n& Objectives", SKY),
    ("2", "\u26A0", "Safety Risk\nManagement", TEAL),
    ("3", "\u2705", "Safety\nAssurance", GOLD),
    ("4", "\U0001F4E2", "Safety\nPromotion", RGBColor(0x9B, 0x59, 0xB6)),
]
ab = AnimBuilder()
ab.add(sid(intro), "fade")
cw = Inches(2.75)
gap = Inches(0.28)
total = 4 * cw + 3 * gap
startx = (EMU_W - total) / 2
cy = Inches(3.55)
ch = Inches(2.7)
for i, (num, g, label, col) in enumerate(comps):
    x = startx + i * (cw + gap)
    card = add_rect(s, x, cy, cw, ch, CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, cy, cw, Inches(0.16), col)
    bdg = add_icon_badge(s, x + (cw - Inches(1.0)) / 2, cy + Inches(0.45), Inches(1.0), col, g, glyph_size=Pt(30))
    add_text(s, x, cy + Inches(0.16), cw, Inches(0.4),
             [[("COMPONENT %s" % num, {"size": Pt(11), "bold": True, "color": col})]],
             align=PP_ALIGN.CENTER)
    lbl = add_text(s, x + Inches(0.15), cy + Inches(1.6), cw - Inches(0.3), Inches(1.0),
                   [[(label.split("\n")[0], {"size": Pt(16), "bold": True, "color": NAVY})],
                    [(label.split("\n")[1], {"size": Pt(16), "bold": True, "color": NAVY})]],
                   align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=Pt(0))
    ab.add(sid(card), "flyBottom", delay=(150 if i else 200))
    ab.add(sid(bdg), "zoom", delay=0)
    ab.add(sid(lbl), "fade", delay=0)
footer(s, 4)
apply_anim(s, ab)


# ============================================================================
# SLIDE 5 — TWELVE KEY ELEMENTS (Part 1)
# ============================================================================
def elements_slide(idx, title_note, blocks, page):
    s = prs.slides.add_slide(BLANK)
    add_solid_bg(s, LIGHT)
    section_header(s, "SMS FRAMEWORK", "The Twelve Key Elements", page)
    add_text(s, Inches(9.5), Inches(0.7), Inches(2.9), Inches(0.4),
             [[(title_note, {"size": Pt(12), "bold": True, "color": TEAL})]], align=PP_ALIGN.RIGHT)
    ab = AnimBuilder()
    y = Inches(1.8)
    for (num, comp_title, col, items) in blocks:
        head = add_rect(s, Inches(0.95), y, Inches(11.45), Inches(0.62), col, MSO_SHAPE.ROUNDED_RECTANGLE)
        htx = add_text(s, Inches(1.25), y, Inches(11), Inches(0.62),
                       [[("%s.  %s" % (num, comp_title), {"size": Pt(18), "bold": True, "color": WHITE})]],
                       anchor=MSO_ANCHOR.MIDDLE)
        ab.add(sid(head), "flyLeft", delay=0)
        ab.add(sid(htx), "fade", delay=0)
        y = y + Inches(0.78)
        for letter, item in items:
            bdg = add_icon_badge(s, Inches(1.35), y + Inches(0.03), Inches(0.4), col, letter,
                                 glyph_color=WHITE, glyph_size=Pt(13))
            itx = add_text(s, Inches(1.95), y, Inches(10.3), Inches(0.42),
                           [[(item, {"size": Pt(16.5), "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
            ab.add(sid(bdg), "zoom", delay=0)
            ab.add(sid(itx), "fade", delay=60)
            y = y + Inches(0.5)
        y = y + Inches(0.14)
    footer(s, page)
    apply_anim(s, ab)


elements_slide(
    1, "Elements 1 \u2013 7",
    [
        ("1", "Safety Policy & Objectives", SKY, [
            ("a", "Management commitment"),
            ("b", "Safety accountability and responsibilities"),
            ("c", "Appointment of key safety personnel"),
            ("d", "Coordination of emergency response planning"),
            ("e", "SMS documentation"),
        ]),
        ("2", "Safety Risk Management", TEAL, [
            ("a", "Hazard identification"),
            ("b", "Safety risk assessment and mitigation"),
        ]),
    ],
    5,
)

# ============================================================================
# SLIDE 6 — TWELVE KEY ELEMENTS (Part 2)
# ============================================================================
elements_slide(
    2, "Elements 8 \u2013 12",
    [
        ("3", "Safety Assurance", GOLD, [
            ("a", "Safety performance monitoring and measurement"),
            ("b", "The management of change"),
            ("c", "Continuous improvement of the SMS"),
        ]),
        ("4", "Safety Promotion", RGBColor(0x9B, 0x59, 0xB6), [
            ("a", "Training and education"),
            ("b", "Safety communication"),
        ]),
    ],
    6,
)


# ============================================================================
# SLIDE 7 — THE WEAKEST LINK
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s, NAVY, DEEP, angle=90)
add_rect(s, 0, EMU_H - Inches(0.14), EMU_W, Inches(0.14), GOLD)
add_rect(s, Inches(0.95), Inches(0.85), Inches(0.16), Inches(0.75), GOLD)
th = add_text(s, Inches(1.3), Inches(0.85), Inches(11), Inches(0.9),
              [[("The Real Challenge Today", {"size": Pt(30), "bold": True, "color": WHITE})]])
bdg1 = add_icon_badge(s, Inches(1.1), Inches(2.35), Inches(0.95), TEAL, "\u2699", glyph_size=Pt(28))
b1 = add_text(s, Inches(2.35), Inches(2.35), Inches(9.9), Inches(1.5),
              [[("Aircraft are now extremely reliable. ", {"size": Pt(21), "bold": True, "color": WHITE}),
                ("The weakest link is no longer the engine \u2014 it\u2019s the system around it: "
                 "procedures, training, communication, fatigue and pressure.",
                 {"size": Pt(21), "color": SKY_LT})]], line_spacing=1.2)
bdg2 = add_icon_badge(s, Inches(1.1), Inches(4.85), Inches(0.95), GOLD, "\U0001F6E1",
                      glyph_color=NAVY, glyph_size=Pt(26))
b2 = add_text(s, Inches(2.35), Inches(4.95), Inches(9.9), Inches(1.2),
              [[("SMS is the tool to manage that system.",
                 {"size": Pt(24), "bold": True, "color": GOLD})]], anchor=MSO_ANCHOR.MIDDLE)

ab = AnimBuilder()
ab.add(sid(th), "flyLeft")
ab.add(sid(bdg1), "zoom", delay=200)
ab.add(sid(b1), "fade", delay=0)
ab.add(sid(bdg2), "zoom", delay=200)
ab.add(sid(b2), "flyLeft", delay=0)
apply_anim(s, ab)


# ============================================================================
# SLIDE 8 — REGULATORY GUIDANCE
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_solid_bg(s, LIGHT)
section_header(s, "COMPLIANCE", "Regulatory Guidance", 8)

regs = [
    ("\U0001F30D", "ICAO Annex 19 / Doc 9859", "Standard Practices \u2014 Safety Management Manual (SMM)", SKY),
    ("\U0001F3DB", "DGCA CAR", "Section 1, Series C, Part I", TEAL),
]
ab = AnimBuilder()
y = Inches(2.4)
for i, (g, title, sub, col) in enumerate(regs):
    card = add_rect(s, Inches(1.6), y, Inches(10.1), Inches(1.55), CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, Inches(1.6), y, Inches(0.16), Inches(1.55), col)
    bdg = add_icon_badge(s, Inches(2.0), y + Inches(0.33), Inches(0.9), col, g, glyph_size=Pt(26))
    tx = add_text(s, Inches(3.3), y + Inches(0.28), Inches(8), Inches(1.0),
                  [[(title, {"size": Pt(22), "bold": True, "color": NAVY})],
                   [(sub, {"size": Pt(15), "color": GRAY})]], line_spacing=1.1, space_after=Pt(2))
    ab.add(sid(card), "flyBottom", delay=(180 if i else 0))
    ab.add(sid(bdg), "zoom", delay=0)
    ab.add(sid(tx), "fade", delay=0)
    y = y + Inches(1.95)
footer(s, 8)
apply_anim(s, ab)


# ============================================================================
# SLIDE 9 — THANK YOU
# ============================================================================
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s, DEEP, NAVY, angle=45)
cc = add_rect(s, Inches(9.8), Inches(3.6), Inches(5), Inches(5), NAVY, MSO_SHAPE.OVAL)
cc.fill.fore_color.rgb = RGBColor(0x16, 0x39, 0x6B)
bdg = add_icon_badge(s, (EMU_W - Inches(1.2)) / 2, Inches(1.9), Inches(1.2), TEAL, "\u2708", glyph_size=Pt(34))
ty = add_text(s, 0, Inches(3.35), EMU_W, Inches(1.2),
              [[("Thank You", {"size": Pt(56), "bold": True, "color": WHITE})]], align=PP_ALIGN.CENTER)
sub = add_text(s, 0, Inches(4.65), EMU_W, Inches(0.6),
               [[("Safety is not an act, it is a habit \u2014 managed, measured and improved every day.",
                  {"size": Pt(17), "italic": True, "color": SKY_LT})]], align=PP_ALIGN.CENTER)
add_rect(s, (EMU_W - Inches(1.6)) / 2, Inches(5.5), Inches(1.6), Inches(0.06), GOLD)
add_rect(s, 0, EMU_H - Inches(0.14), EMU_W, Inches(0.14), GOLD)

ab = AnimBuilder()
ab.add(sid(bdg), "zoom")
ab.add(sid(ty), "fade", delay=200)
ab.add(sid(sub), "flyBottom", delay=200)
apply_anim(s, ab)


# ----------------------------------------------------------------------------
out = "Safety_Management_System.pptx"
prs.save(out)
print("Saved:", out, "| slides:", len(prs.slides._sldIdLst))
